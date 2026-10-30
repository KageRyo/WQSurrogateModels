"""Export completed edge CSV statistics without changing scientific calculations.

Requires openpyxl. Huge per-window stress data stays in its source CSV, indexed
with a hash and row count. No derived statistics or spreadsheet formulas are added.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SHEETS = {
    "metrics_summary": "Metrics_summary", "metrics_by_seed": "Metrics_by_seed",
    "bootstrap_ci": "Bootstrap_CI", "paired_error_tests": "Paired_tests",
    "best_by_experiment": "Best_by_experiment", "cpu_inference_timing": "CPU_timing",
    "component_resources": "Component_resources", "error_by_wqi_band_by_seed": "WQI_band_errors",
    "prediction_parity": "Prediction_parity", "global_stress": "Global_stress",
    "stress107_detection_summary": "Stress107_detection",
    "stress107_severity_monotonicity": "Stress107_monotonicity",
}
INTEGER = re.compile(r"^-?(?:0|[1-9][0-9]*)$")
NUMBER = re.compile(r"^-?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?$")
FONT = Font(name="Arial", size=10)
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")


def digest(path):
    checksum = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            checksum.update(block)
    return checksum.hexdigest()


def value(raw):
    if raw == "":
        return None
    if raw in {"True", "False"}:
        return raw == "True"
    if INTEGER.fullmatch(raw) and len(raw.lstrip("-")) <= 15:
        return int(raw)
    if NUMBER.fullmatch(raw):
        return float(raw)
    return raw


def format_sheet(sheet):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 28
    for row in sheet:
        for cell in row:
            cell.font = FONT
            if isinstance(cell.value, float):
                cell.number_format = "0.000000E+00" if 0 < abs(cell.value) < 0.0001 else "0.000000;[Red]-0.000000;0"
            elif isinstance(cell.value, int) and not isinstance(cell.value, bool):
                cell.number_format = "0"
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for column in sheet.columns:
        width = min(55, max(15, max(len(str(cell.value or "")) for cell in list(column)[:100]) + 2))
        sheet.column_dimensions[get_column_letter(column[0].column)].width = width


def export(source, output):
    source, output = Path(source).resolve(), Path(output).resolve()
    if output.exists():
        raise FileExistsError(output)
    manifest = json.loads((source / "manifest.json").read_text())
    book = Workbook()
    readme = book.active
    readme.title = "Readme"
    readme.append(["項目", "說明／來源值"])
    notes = [
        ("資料來源", str(source)),
        ("資料時間 UTC", manifest["created_utc"]),
        ("完成工作數（manifest 原值）", manifest["verified_jobs"]),
        ("警示：all_parity_pass", manifest["all_parity_pass"]),
        ("Parity 判讀", "全部失敗工作逐項列於 Parity_findings；不是全數預測一致。Prediction_parity 保留通過及失敗的完整 frame 列。"),
        ("數值語義", "既有 CSV 科學統計原值，不重新計算、不裁切資料、不新增推論或排名。"),
        ("時間單位與範圍", manifest["timing_contract"]),
        ("記憶體單位與範圍", manifest["memory_contract"]),
        ("訓練範圍", manifest["training_contract"]),
        ("單位", "*_seconds、*_s 為秒；*_ms 為毫秒；*_mib 為 MiB；分數／誤差沿用 WQI5。"),
        ("比例", "accuracy、macro_f1、*_rate、pct_* 為 0–1 比例；維持原數值格式，沒有乘以 100。p 值亦為原值。"),
        ("不內嵌的大型資料", "stress107_windows.csv 完整保留 CSV，不抽樣；Source_index 列明列數、相對位置與 SHA256。"),
        ("完整性", "Source_index 每列是一個來源檔；embedded_rows 為逐筆讀取來源後記錄的匯入列數（不含標題）。未內嵌者為 0。"),
        ("開啟方式", "唯讀結果工作簿；可篩選欄位。完整 Stress107 CSV 與本檔需一起保留。來源路徑相對於工作簿所在資料夾。"),
        ("Parity 差異上限（manifest 原值）", manifest["max_prediction_difference"]),
    ]
    for row in notes:
        readme.append(row)
    findings = book.create_sheet("Parity_findings")
    findings.append(["job_id（manifest parity_difference_jobs 原值）"])
    for job in manifest["parity_difference_jobs"]:
        findings.append([job])
    index = book.create_sheet("Source_index")
    index.append(["source_file", "workbook_sheet", "source_rows", "embedded_rows", "bytes", "sha256", "relative_source_path", "scope"])
    expected = {}
    failed_jobs = set()
    csv_paths = sorted(source.glob("*.csv"))
    if {p.stem for p in csv_paths} != {*SHEETS, "stress107_windows"}:
        raise ValueError("CSV inventory changed; explicitly classify every new source")
    import os
    for path in csv_paths:
        sheet_name = SHEETS.get(path.stem)
        sheet = book.create_sheet(sheet_name) if sheet_name else None
        count = 0
        with path.open(newline="", encoding="utf-8") as stream:
            reader = csv.reader(stream)
            headers = next(reader)
            if sheet:
                sheet.append(headers)
            for raw in reader:
                if len(raw) != len(headers):
                    raise ValueError(f"Malformed CSV row: {path}:{count + 2}")
                count += 1
                if sheet:
                    sheet.append([value(item) for item in raw])
                if path.stem == "prediction_parity" and raw[headers.index("pass")] == "False":
                    failed_jobs.add(raw[headers.index("job_id")])
        if sheet:
            expected[sheet_name] = count
            format_sheet(sheet)
        index.append([path.name, sheet_name or "EXTERNAL CSV — all rows retained", count,
                      count if sheet else 0, path.stat().st_size, digest(path),
                      os.path.relpath(path, output.parent), "complete" if sheet else "not embedded; no truncation/sample"])
        index.cell(index.max_row, 7).hyperlink = os.path.relpath(path, output.parent)
    index.append(["manifest.json", "Readme + Parity_findings", None, None,
                  (source / "manifest.json").stat().st_size, digest(source / "manifest.json"),
                  os.path.relpath(source / "manifest.json", output.parent), "metadata and full failed-job list"])
    if failed_jobs != set(manifest["parity_difference_jobs"]):
        raise ValueError("CSV parity failures disagree with manifest")
    for sheet in (readme, findings, index):
        format_sheet(sheet)
    readme.column_dimensions["A"].width = 35
    readme.column_dimensions["B"].width = 110
    for row in readme.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        readme.row_dimensions[row[0].row].height = 44
    readme["B5"].fill = PatternFill("solid", fgColor="FFC7CE")
    findings.column_dimensions["A"].width = 105
    for row in findings.iter_rows(min_row=2):
        row[0].fill = PatternFill("solid", fgColor="FFC7CE")
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    book.close()
    check = load_workbook(output, read_only=True, data_only=False)
    for name, count in expected.items():
        if check[name].max_row != count + 1:
            raise ValueError(f"Reload row mismatch: {name}")
    errors, formulas = [], []
    for sheet in check:
        for row in sheet:
            for cell in row:
                if cell.data_type == "e":
                    errors.append(f"{sheet.title}!{cell.coordinate}")
                if cell.data_type == "f":
                    formulas.append(f"{sheet.title}!{cell.coordinate}")
    if errors or formulas:
        raise ValueError(f"Unexpected cell errors/formulas: {errors}, {formulas}")
    # Compare every embedded value with its CSV, allowing only XLSX numeric
    # serialization precision. This catches shifted columns or altered types.
    for path in csv_paths:
        if path.stem not in SHEETS:
            continue
        with path.open(newline="", encoding="utf-8") as stream:
            source_rows = csv.reader(stream)
            headers = next(source_rows)
            sheet = check[SHEETS[path.stem]]
            if tuple(headers) != next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)):
                raise ValueError(f"Reloaded headers differ: {path.name}")
            for number, (raw, actual) in enumerate(zip(source_rows, sheet.iter_rows(min_row=2, values_only=True)), 2):
                for column, (text, cell) in enumerate(zip(raw, actual), 1):
                    original = value(text)
                    equal = (math.isclose(original, cell, rel_tol=1e-14, abs_tol=0)
                             if isinstance(original, float) and isinstance(cell, (float, int))
                             else original == cell)
                    if not equal:
                        raise ValueError(f"Reloaded value differs: {path.name}:{number}:{column}")
    reloaded_jobs = {row[0] for row in check["Parity_findings"].iter_rows(min_row=2, values_only=True)}
    if reloaded_jobs != failed_jobs:
        raise ValueError("Reloaded parity list differs")
    check.close()
    print(json.dumps({"output": str(output), "bytes": output.stat().st_size,
                      "verified_jobs": manifest["verified_jobs"], "failed_parity_jobs": len(failed_jobs),
                      "sheet_rows": expected, "cell_errors": 0, "formulas": 0}, ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    export(args.source, args.output)


if __name__ == "__main__":
    main()
